# -*- coding: utf-8 -*-
"""导出合作者标注包素材：原文 txt + Excel 标注模板（带枚举下拉框）。

用法: python src/export_annotation_pack.py
输入: data/annotation/seed_v1.jsonl
输出: 02_标注与复核/标注包/原文/doc_<id>.txt 与 02_标注与复核/标注包/事件标注模板.xlsx
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SEED = ROOT / "data" / "annotation" / "seed_v1.jsonl"
PACK = ROOT / "02_标注与复核" / "标注包"
TEXTS = PACK / "原文"

GREEN = PatternFill("solid", fgColor="1A4D2E")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
WRAP = Alignment(wrap_text=True, vertical="top")

docs = [json.loads(l) for l in SEED.read_text(encoding="utf-8").strip().splitlines()]

TEXTS.mkdir(parents=True, exist_ok=True)
for d in docs:
    (TEXTS / f"doc_{d['doc_id']}.txt").write_text(
        f"标题：{d['title']}\n发布日期：{d['publish_date']}\n来源：{d['source']}\n\n{d['content_text']}",
        encoding="utf-8",
    )

wb = Workbook()

def style_header(ws, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(1, i, h)
        c.fill = GREEN
        c.font = HDR_FONT
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

# ---------- Sheet1 填写说明 ----------
ws = wb.active
ws.title = "填写说明"
ws.column_dimensions["A"].width = 110
lines = [
    "绿证市场政策事件标注模板 v0.1",
    "",
    "【工作流程】",
    "1. 打开「文档列表」表，按 doc_id 找到「原文」文件夹里对应的 doc_<id>.txt，通读全文",
    "2. 判断文中有哪些「会改变绿证市场供需、价格或规则的事实」，每个事实 = 一条事件",
    "3. 到「事件标注」表填写：每条事件一行，doc_id 填所属文档编号",
    "4. 一篇文档可以有 0 条、1 条或多条事件；0 条时无需在「事件标注」表填任何内容",
    "5. 标完一篇，在「文档列表」表最后一列打勾并填标注人姓名",
    "",
    "【事件类型速查】（详细定义与例子见《标注指南.pdf》第 2 节）",
    "A 政策发布（A1顶层规划 A2实施细则 A3管理办法 A4征求意见）",
    "B 考核范围扩容（B1新增考核行业 B2新增考核主体 B3考核力度升级）",
    "C 消纳责任权重调整（C1年度权重下达 C2权重上调/下调 C3省级分解方案）",
    "D 价格信号（D1官方价格通报 D2成交量通报 D3价格异常说明）",
    "E 供需冲击（E1核发放量 E2需求放量 E3供给受限）",
    "F 市场规则变更（F1交易规则 F2绿电直连 F3核销与有效期 F4国际互认）",
    "",
    "【关键规则】",
    "1. evidence_span 必须是原文原句复制粘贴（≤100字），不许改写——无证据的事件视为无效",
    "2. effective_date 填政策管辖/生效时间（不是文件印发时间），精度不得超过原文：原文只到年份就填 2025，不要编 2025-01-01；没有就填「未明确」",
    "3. strength 看措辞：约束性指标/考核/应 → 约束性；预期/力争 → 预期性；鼓励/支持 → 倡导性",
    "4. 新闻评论类文章只标其中引用的官方事实，记者观点不标",
    "5. 同一事件同子类同对象只保留一条；但力度或年份不同的（如2025约束性 vs 2026预期性）要拆成两条",
    "6. price_direction_prior 填你作为领域研究者对该事件影响绿证价格方向的判断（上行/下行/中性/不明）",
    "",
    "【示例】见「示例-doc42」表：对《2025年消纳责任权重通知》的完整标注",
    "【疑问】拿不准的先在「备注」列写疑问，不要猜；汇总后与课题负责人统一仲裁",
]
for i, t in enumerate(lines, 1):
    c = ws.cell(i, 1, t)
    if i == 1:
        c.font = Font(bold=True, size=13, color="1A4D2E")

# ---------- Sheet2 文档列表 ----------
ws2 = wb.create_sheet("文档列表")
style_header(ws2, ["doc_id", "来源", "标题", "发布日期", "命中关键词", "原文文件", "已标注(打勾)", "标注人"],
             [8, 8, 55, 11, 18, 15, 11, 9])
for i, d in enumerate(docs, 2):
    ws2.cell(i, 1, d["doc_id"])
    ws2.cell(i, 2, d["source"])
    c = ws2.cell(i, 3, d["title"]); c.alignment = WRAP
    ws2.cell(i, 4, d["publish_date"])
    ws2.cell(i, 5, d["keywords_matched"])
    ws2.cell(i, 6, f"doc_{d['doc_id']}.txt")

# ---------- Sheet3 事件标注 ----------
ws3 = wb.create_sheet("事件标注")
headers = ["doc_id", "event_type", "subtype", "actor", "object", "strength",
           "effective_date", "region", "industry", "price_direction_prior", "evidence_span", "备注"]
style_header(ws3, headers, [8, 11, 9, 22, 22, 10, 13, 9, 20, 18, 50, 18])
for i in range(2, 402):
    ws3.cell(i, 1).alignment = WRAP
    ws3.cell(i, 11).alignment = WRAP

dv_type = DataValidation(type="list", formula1='"A,B,C,D,E,F"', allow_blank=True)
dv_sub = DataValidation(type="list", formula1='"A1,A2,A3,A4,B1,B2,B3,C1,C2,C3,D1,D2,D3,E1,E2,E3,F1,F2,F3,F4"', allow_blank=True)
dv_str = DataValidation(type="list", formula1='"约束性,预期性,倡导性"', allow_blank=True)
dv_dir = DataValidation(type="list", formula1='"上行,下行,中性,不明"', allow_blank=True)
for dv, col in [(dv_type, "B"), (dv_sub, "C"), (dv_str, "F"), (dv_dir, "J")]:
    ws3.add_data_validation(dv)
    dv.add(f"{col}2:{col}400")

# ---------- Sheet4 示例 ----------
ws4 = wb.create_sheet("示例-doc42")
ws4.column_dimensions["A"].width = 110
example = [
    "示例文档：doc_42《关于2025年可再生能源电力消纳责任权重及有关事项的通知》（发改办能源〔2025〕669号）",
    "该文档可标出 4 条事件，在「事件标注」表中应填为以下 4 行：",
    "",
    "行1: doc_id=42 | event_type=C | subtype=C1 | actor=国家发展改革委、国家能源局 | object=各省（自治区、直辖市） | strength=约束性 | effective_date=2025 | region=全国 | industry=(空) | price_direction_prior=上行 | evidence_span=本次下达的2025年可再生能源电力消纳责任权重为约束性指标，按此对各省（自治区、直辖市）进行考核评估",
    "行2: doc_id=42 | C | C1 | actor同上 | object同上 | strength=预期性 | effective_date=2026 | region=全国 | 上行 | evidence_span=2026年权重为预期性指标，各省（自治区、直辖市）要按此做好项目储备",
    "行3: doc_id=42 | B | B1 | actor同上 | object=钢铁、水泥、多晶硅行业和新建数据中心 | strength=约束性 | effective_date=2025 | region=全国 | industry=钢铁,水泥,多晶硅,数据中心 | 上行 | evidence_span=在电解铝行业基础上，2025年增设钢铁、水泥、多晶硅行业和国家枢纽节点新建数据中心绿色电力消费比例",
    "行4: doc_id=42 | F | F3 | actor同上 | object=消纳责任权重核算 | strength=约束性 | effective_date=2025 | region=全国 | 上行 | evidence_span=2025年可再生能源电力消纳责任权重应在当年完成，不再转移至2026年",
    "",
    "注意行1与行2：同一份文件的2025年权重（约束性）和2026年权重（预期性）必须拆成两条事件，因为约束力度不同、对价格的含义不同。",
]
for i, t in enumerate(example, 1):
    c = ws4.cell(i, 1, t)
    c.alignment = WRAP
    if i == 1:
        c.font = Font(bold=True, size=11, color="1A4D2E")

out = PACK / "事件标注模板.xlsx"
wb.save(out)
print(f"原文导出 {len(docs)} 篇 -> {TEXTS}")
print(f"Excel 模板 -> {out}")
