# -*- coding: utf-8 -*-
"""构建 gold set v1：从人工复核 Excel（02_标注与复核/标注包/WWW绿证政策事件复核.xlsx）生成
data/annotation/gold_v1.jsonl（60 篇全覆盖）+ pending_review.json（存疑/待人工确认）。

裁决规则：
  采K3 → K3版本；采GPT6 → GPT6版本；修改 → 修正后值；删除 → 排除；
  存疑 → pending_review.json
  一致事件「确认」→ 直接入 gold（object/industry 若为 "K3: x / GPT6: y" 合并格式取更长侧，
  并记录）；一致事件「修改」（备注为自由文本）→ 不机器解析，入 pending_review.json。
每条 gold 事件用 common.validate 校验，问题逐条列出，不静默修复。
"""
import json
import logging
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import validate  # noqa: E402

logger = logging.getLogger("build_gold_v1")
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s",
                    stream=sys.stdout)

XLSX = ROOT / "02_标注与复核" / "标注包" / "WWW绿证政策事件复核.xlsx"
SEED = ROOT / "data" / "annotation" / "seed_v1.jsonl"
GOLD_OUT = ROOT / "data" / "annotation" / "gold_v1.jsonl"
PENDING_OUT = ROOT / "data" / "annotation" / "pending_review.json"

VERSION_FIELDS = ["subtype", "actor", "object", "strength", "effective_date",
                  "region", "industry", "price_direction_prior", "evidence_span"]
DATE_RE = __import__("re").compile(r"^(\d{4}(-\d{2}(-\d{2})?)?|未明确)$")


def parse_version(text):
    """解析 'subtype=A1 | actor=... | ...' 版本串为事件 dict。"""
    ev = {}
    for seg in str(text).split(" | "):
        if "=" not in seg:
            continue
        k, v = seg.split("=", 1)
        k, v = k.strip(), v.strip()
        if k in VERSION_FIELDS:
            ev[k] = v
    if "subtype" in ev:
        ev["event_type"] = ev["subtype"][0]
    ind = ev.get("industry", "")
    ev["industry"] = [x.strip() for x in ind.split(",") if x.strip()] if ind else []
    return ev


def pick_longer_side(cell):
    """一致表措辞字段若为 'K3: xxx / GPT6: yyy' 合并格式，取更长侧。"""
    s = str(cell or "")
    if s.startswith("K3: ") and " / GPT6: " in s:
        a, b = s[4:].split(" / GPT6: ", 1)
        pick = a if len(a) >= len(b) else b
        return ("" if pick == "∅" else pick), f"合并措辞取更长侧({'K3' if len(a) >= len(b) else 'GPT6'})"
    return s, None


def main():
    seed = [json.loads(l) for l in SEED.read_text(encoding="utf-8").splitlines() if l.strip()]
    seed_by_id = {d["doc_id"]: d for d in seed}
    wb = load_workbook(XLSX)

    gold_events = {d["doc_id"]: [] for d in seed}
    pending = []
    problems = []
    adoption = {"采K3": 0, "采GPT6": 0, "修改": 0, "一致确认": 0}
    merge_notes = []

    # ---- 分歧事件 ----
    ws = wb["分歧事件"]
    header = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        doc_id, subtype, dtype = row[col["doc_id"]], row[col["subtype"]], row[col["分歧类型"]]
        verdict = row[col["人工裁决"]]
        ctx = {"doc_id": doc_id, "subtype": subtype, "分歧类型": dtype,
               "差异字段": row[col["差异字段"]],
               "K3版本": row[col["K3版本"]], "GPT6版本": row[col["GPT6版本"]],
               "evidence_span_K3": row[col["evidence_span_K3"]],
               "evidence_span_GPT6": row[col["evidence_span_GPT6"]],
               "人工裁决": verdict, "修正后值": row[col["修正后值"]],
               "备注": row[col["备注"]]}
        if verdict == "删除":
            continue
        if verdict == "存疑":
            pending.append({"来源": "分歧事件", **ctx})
            continue
        if verdict == "采K3":
            ev = parse_version(row[col["K3版本"]])
            # 版本列只含 8 个字段，evidence 在独立列：采纳即采纳该侧完整版本，补 evidence
            if not ev.get("evidence_span"):
                ev["evidence_span"] = row[col["evidence_span_K3"]] or ""
        elif verdict == "采GPT6":
            ev = parse_version(row[col["GPT6版本"]])
            if not ev.get("evidence_span"):
                ev["evidence_span"] = row[col["evidence_span_GPT6"]] or ""
        elif verdict == "修改":
            ev = parse_version(row[col["修正后值"]])
            if not ev.get("evidence_span"):
                # 人工只改了字段值、未重抄证据（50 行中 35 行如此）。补引规则（显式）：
                # K3/GPT6 证据一致或互为子串 → 补引较长者并记录；否则留空进问题清单。
                ek = (row[col["evidence_span_K3"]] or "").strip()
                eg = (row[col["evidence_span_GPT6"]] or "").strip()
                if ek and eg and (ek in eg or eg in ek):
                    ev["evidence_span"] = ek if len(ek) >= len(eg) else eg
                    merge_notes.append(
                        f"doc {doc_id} {subtype}: 修改行证据补引（两侧一致/包含，取较长）")
                elif ek and not eg:
                    ev["evidence_span"] = ek
                    merge_notes.append(f"doc {doc_id} {subtype}: 修改行证据补引（仅K3侧有）")
                elif eg and not ek:
                    ev["evidence_span"] = eg
                    merge_notes.append(f"doc {doc_id} {subtype}: 修改行证据补引（仅GPT6侧有）")
                else:
                    problems.append(
                        f"doc {doc_id} {subtype}: 人工「修改」未含 evidence_span 且两侧证据不一致，"
                        f"留空待人工补（不静默修复）")
        else:
            problems.append(f"doc {doc_id} {subtype}: 未识别裁决 '{verdict}'")
            continue
        adoption[verdict] += 1
        gold_events[doc_id].append(ev)

    # ---- 一致事件 ----
    ws1 = wb["一致事件"]
    header1 = [c.value for c in ws1[1]]
    c1 = {name: i for i, name in enumerate(header1)}
    for row in ws1.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        doc_id = row[c1["doc_id"]]
        confirm = row[c1["人工确认"]]
        if confirm == "确认":
            ev = {f: row[c1[f]] for f in
                  ["subtype", "actor", "strength", "effective_date", "region",
                   "price_direction_prior", "evidence_span"]}
            ev["event_type"] = ev["subtype"][0]
            for f in ("object", "industry"):
                v, note = pick_longer_side(row[c1[f]])
                if f == "industry":
                    ev[f] = [x.strip() for x in v.split(",") if x.strip()] if v else []
                else:
                    ev[f] = v
                if note:
                    merge_notes.append(f"doc {doc_id} {f}: {note}")
            adoption["一致确认"] += 1
            gold_events[doc_id].append(ev)
        elif confirm == "修改":
            pending.append({"来源": "一致事件",
                            "说明": "人工确认=修改，备注为自由文本修正说明，机器不解析，需人工判读",
                            "doc_id": doc_id, "subtype": row[c1["subtype"]],
                            "原始行": {f: row[c1[f]] for f in
                                       ["actor", "object", "strength", "effective_date",
                                        "region", "industry", "price_direction_prior",
                                        "evidence_span"]},
                            "备注": row[c1["备注"]]})
        else:
            problems.append(f"一致表 doc {doc_id}: 未处理确认值 '{confirm}'")

    # ---- 校验 + 输出 ----
    n_events = 0
    type_dist = {}
    with GOLD_OUT.open("w", encoding="utf-8") as f:
        for doc in seed:
            doc_id = doc["doc_id"]
            events = gold_events[doc_id]
            n_events += len(events)
            for i, ev in enumerate(events):
                _, errs = validate(json.dumps({"events": [ev]}, ensure_ascii=False))
                ed = ev.get("effective_date", "")
                if ed and not DATE_RE.match(ed):
                    errs.append(f"effective_date 格式异常: {ed}")
                if not ev.get("subtype"):
                    errs.append("缺 subtype")
                for e in errs:
                    problems.append(f"doc {doc_id} 事件{i}({ev.get('subtype')}): {e}")
                if ev.get("event_type"):
                    type_dist[ev["event_type"]] = type_dist.get(ev["event_type"], 0) + 1
            f.write(json.dumps({
                "doc_id": doc_id, "source": doc["source"], "title": doc["title"],
                "publish_date": doc["publish_date"], "ontology_version": "0.1",
                "annotator": "human_adjudicated", "events": events,
            }, ensure_ascii=False) + "\n")

    PENDING_OUT.write_text(json.dumps(pending, ensure_ascii=False, indent=1), encoding="utf-8")

    n_with = sum(1 for v in gold_events.values() if v)
    print("\n===== gold set v1 统计 =====")
    print(f"文档: 60 篇（有事件 {n_with}，无事件 {60 - n_with}）")
    print(f"gold 总事件数: {n_events}")
    print(f"按大类分布: {dict(sorted(type_dist.items()))}")
    total_adopt = sum(adoption.values())
    print(f"采纳构成（共 {total_adopt} 条入 gold）: "
          + ", ".join(f"{k} {v} ({v/total_adopt*100:.1f}%)" for k, v in adoption.items()))
    print(f"存疑/待人工确认: {len(pending)} 条 -> {PENDING_OUT.name}")
    if merge_notes:
        print("一致表合并措辞处理:", merge_notes)
    print("\n===== 校验问题清单 =====")
    if problems:
        for p in problems:
            print(" !", p)
    else:
        print("（空）")
    logger.info("gold 写出: %s；pending: %s", GOLD_OUT, PENDING_OUT)


if __name__ == "__main__":
    main()
